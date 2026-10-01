"""
VitaLens — University Project Presentation Generator
=====================================================
Author: Uzair Shaikh
Faculty: Prof. Swati D. Londhe  |  HOD: Prof. Monica Ma'am

Run:  pip install python-pptx
      python generate_presentation.py

Output: VitaLens_University_Presentation.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import datetime

# ─────────────────────────────────────────────
#  DESIGN TOKENS  (Light Green + Light Blue + White theme)
# ─────────────────────────────────────────────
BG          = RGBColor(0xF0, 0xFB, 0xF4)   # slide background — very light mint
CARD        = RGBColor(0xD8, 0xF3, 0xE3)   # card / panel fill — soft mint green
GREEN       = RGBColor(0x2A, 0xA8, 0x6E)   # primary accent — medium green
BLUE        = RGBColor(0x14, 0x8F, 0xC1)   # secondary accent — sky blue
L_BLUE      = RGBColor(0x90, 0xD5, 0xE4)   # light blue highlight
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
DARK        = RGBColor(0x1A, 0x3C, 0x34)   # dark forest green — body text
MID         = RGBColor(0x3A, 0x7D, 0x6B)   # muted teal-green — secondary text
BORDER      = RGBColor(0xB2, 0xDF, 0xC8)   # card border — pale green

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


# ════════════════════════════════════════════════════════════════════
#  HELPERS
# ════════════════════════════════════════════════════════════════════

def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

def rect(slide, l, t, w, h, color, line=None, lw=0):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    return s

def txt(slide, text, l, t, w, h, size=14, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return tb

def bullets(slide, items, l, t, w, h, size=13, color=DARK, gap=5):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(gap)
        r = p.add_run()
        r.text = f"  {item}"
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"

def top_bar(slide, color=GREEN):
    rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.12), color)

def bot_bar(slide):
    rect(slide, Inches(0), Inches(7.25), SLIDE_W, Inches(0.25), GREEN)
    txt(slide, "VitaLens  |  React Native Mobile App  |  University Project Presentation",
        Inches(0.4), Inches(7.25), Inches(12.5), Inches(0.25),
        size=9, color=WHITE, align=PP_ALIGN.CENTER)

def slide_header(slide, label, title):
    top_bar(slide)
    txt(slide, label.upper(), Inches(0.5), Inches(0.2), Inches(5), Inches(0.3),
        size=9, bold=True, color=GREEN)
    txt(slide, title, Inches(0.5), Inches(0.45), Inches(12.3), Inches(0.7),
        size=28, bold=True, color=DARK)
    rect(slide, Inches(0.5), Inches(1.2), Inches(12.3), Pt(2), GREEN)
    bot_bar(slide)

def card(slide, l, t, w, h):
    rect(slide, l, t, w, h, CARD, BORDER, 0.8)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 1 — TITLE / COVER
# ════════════════════════════════════════════════════════════════════

def slide_01_title():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)

    # Top green header band
    rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.55), GREEN)

    # University label
    txt(slide, "UNIVERSITY PROJECT PRESENTATION",
        Inches(0.5), Inches(0.08), Inches(12.3), Inches(0.4),
        size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Project title on green band
    txt(slide, "VitaLens",
        Inches(0.5), Inches(0.45), Inches(12.3), Inches(0.85),
        size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Subtitle below the band
    txt(slide,
        "A Cross-Platform Health Report & Doctor Booking Mobile Application",
        Inches(0.5), Inches(1.65), Inches(12.3), Inches(0.55),
        size=18, bold=False, color=DARK, align=PP_ALIGN.CENTER)

    # Thin divider
    rect(slide, Inches(3.5), Inches(2.28), Inches(6.3), Pt(2), GREEN)

    # Built with tag
    txt(slide, "Built with  React Native (Expo)  ·  TypeScript  ·  FastAPI",
        Inches(0.5), Inches(2.38), Inches(12.3), Inches(0.4),
        size=13, italic=True, color=MID, align=PP_ALIGN.CENTER)

    # ── Two-column info cards ──
    # Left card — Project Author
    card(slide, Inches(0.5), Inches(2.95), Inches(5.9), Inches(3.5))
    txt(slide, "Project Author",
        Inches(0.7), Inches(3.1), Inches(5.5), Inches(0.4),
        size=13, bold=True, color=GREEN)
    rect(slide, Inches(0.7), Inches(3.52), Inches(5.5), Pt(1.5), BORDER)

    members = [
        ("Uzair Shaikh", "Enrollment No:  22004500210246"),
    ]
    y = Inches(3.65)
    for name, enroll in members:
        txt(slide, name, Inches(0.75), y, Inches(5.4), Inches(0.35),
            size=14, bold=True, color=DARK)
        txt(slide, enroll, Inches(0.75), y + Inches(0.32), Inches(5.4), Inches(0.3),
            size=11, color=MID)
        y += Inches(0.9)

    # Right card — Faculty & Course Details
    card(slide, Inches(6.9), Inches(2.95), Inches(5.95), Inches(3.5))
    txt(slide, "Faculty & Course Details",
        Inches(7.1), Inches(3.1), Inches(5.6), Inches(0.4),
        size=13, bold=True, color=BLUE)
    rect(slide, Inches(7.1), Inches(3.52), Inches(5.6), Pt(1.5), BORDER)

    details = [
        ("Faculty / Guide:", "Prof. Swati D. Londhe"),
        ("HOD:", "Prof. Monica Ma'am"),
        ("Technology:", "React Native Mobile App"),
        ("Academic Year:", "2025 – 2026"),
    ]
    y = Inches(3.65)
    for label, value in details:
        txt(slide, label, Inches(7.15), y, Inches(2.0), Inches(0.32),
            size=11, bold=True, color=MID)
        txt(slide, value, Inches(9.25), y, Inches(2.5), Inches(0.32),
            size=12, bold=False, color=DARK)
        y += Inches(0.75)

    # Bottom green footer bar
    rect(slide, Inches(0), Inches(7.25), SLIDE_W, Inches(0.25), GREEN)
    txt(slide, f"Submitted  —  {datetime.date.today().strftime('%B %Y')}",
        Inches(0.4), Inches(7.25), Inches(12.5), Inches(0.25),
        size=9, color=WHITE, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 2 — PROBLEM STATEMENT & OBJECTIVE
# ════════════════════════════════════════════════════════════════════

def slide_02_problem():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)
    slide_header(slide, "Project Background", "Problem Statement & Objectives")

    # Problem box
    card(slide, Inches(0.5), Inches(1.35), Inches(5.9), Inches(5.55))
    txt(slide, "Problem Statement",
        Inches(0.7), Inches(1.5), Inches(5.5), Inches(0.4),
        size=15, bold=True, color=GREEN)

    problems = [
        "Patients receive complex lab reports in dense medical jargon — difficult to understand without a doctor",
        "Finding the right specialist doctor is confusing and time-consuming with no digital guidance",
        "Lab results on paper get lost; there is no organized, digital way for patients to track health data over time",
        "Booking doctor appointments still relies on phone calls or walk-ins — no convenient mobile solution",
        "No centralized platform connects report review, symptom logging, and consultation booking",
    ]
    bullets(slide, problems, Inches(0.7), Inches(2.0), Inches(5.5), Inches(4.5),
            size=12, color=DARK, gap=6)

    # Objectives box
    card(slide, Inches(6.9), Inches(1.35), Inches(5.95), Inches(5.55))
    txt(slide, "Project Objectives",
        Inches(7.1), Inches(1.5), Inches(5.6), Inches(0.4),
        size=15, bold=True, color=BLUE)

    objectives = [
        "Build a fully functional React Native mobile app for Android & iOS using Expo",
        "Allow users to upload medical lab reports and view extracted biomarker data",
        "Provide plain-English summaries of report results for non-medical users",
        "Enable symptom logging and suggest the appropriate medical specialty",
        "Display a searchable directory of doctors with real profiles and availability",
        "Allow users to book, manage, and track appointments end-to-end",
        "Implement secure login, user profiles, and privacy consent controls",
    ]
    bullets(slide, objectives, Inches(7.1), Inches(2.0), Inches(5.6), Inches(4.5),
            size=12, color=DARK, gap=5)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 3 — TECHNOLOGY STACK
# ════════════════════════════════════════════════════════════════════

def slide_03_tech_stack():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)
    slide_header(slide, "Technical Overview", "Technology Stack Used")

    col_data = [
        {
            "icon": "📱",
            "title": "Frontend  —  Mobile App",
            "color": GREEN,
            "items": [
                "React Native with Expo SDK",
                "TypeScript (full type safety)",
                "NativeWind (Tailwind CSS for RN)",
                "React Navigation (Stack + Bottom Tabs)",
                "Zustand (global state management)",
                "Axios (API calls with JWT auth)",
                "Lucide React Native (icons)",
                "Expo ImagePicker (file uploads)",
            ],
        },
        {
            "icon": "⚙️",
            "title": "Backend  —  API Server",
            "color": BLUE,
            "items": [
                "FastAPI (Python 3.12)",
                "SQLAlchemy 2.0 Async ORM",
                "PostgreSQL (relational database)",
                "Alembic (DB migrations)",
                "JWT Authentication (Bearer tokens)",
                "Firebase Admin SDK (notifications)",
                "Uvicorn (ASGI server)",
                "Pytest (automated testing)",
            ],
        },
        {
            "icon": "🛠️",
            "title": "Dev Tools & Extras",
            "color": MID,
            "items": [
                "Expo Go (live preview on device)",
                "Metro Bundler (JS bundler)",
                "Swagger UI (API documentation)",
                "Git (version control)",
                "VS Code (IDE)",
                "Postman (API testing)",
                "ESLint + TypeScript compiler",
                "React Admin Web (admin panel)",
            ],
        },
    ]

    col_x = [Inches(0.5), Inches(4.95), Inches(9.4)]
    for ci, col in enumerate(col_data):
        cx = col_x[ci]
        # Header strip
        rect(slide, cx, Inches(1.35), Inches(4.0), Inches(0.55), col["color"])
        txt(slide, f"{col['icon']}  {col['title']}",
            cx + Inches(0.12), Inches(1.42), Inches(3.76), Inches(0.42),
            size=13, bold=True, color=WHITE)
        # Body card
        card(slide, cx, Inches(1.9), Inches(4.0), Inches(5.05))
        bullets(slide, col["items"],
                cx + Inches(0.15), Inches(2.0), Inches(3.7), Inches(4.8),
                size=12, color=DARK, gap=5)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 4 — KEY FEATURES (Module 1 — Reports & Home)
# ════════════════════════════════════════════════════════════════════

def slide_04_features_reports():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)
    slide_header(slide, "Key Features  |  Module 1", "Home Dashboard & Report Management")

    feature_data = [
        {
            "icon": "🏠",
            "title": "Home Dashboard",
            "points": [
                "Personalized greeting with user name",
                "Quick action buttons for all modules",
                "At-a-glance health summary cards",
                "Notification access from the dashboard",
            ],
        },
        {
            "icon": "📤",
            "title": "Upload Lab Report",
            "points": [
                "Upload PDF or image from camera/gallery",
                "Progress indicator during file upload",
                "Report stored securely per user account",
                "Status tracking: Pending → Processing → Done",
            ],
        },
        {
            "icon": "🔬",
            "title": "Report Analysis View",
            "points": [
                "Biomarker cards with Normal / High / Low flags",
                "Plain-English summary of results",
                "Medical terminology glossary modal",
                "Clinical disclaimer clearly displayed",
            ],
        },
        {
            "icon": "📈",
            "title": "Biomarker Trends",
            "points": [
                "Time-series charts per biomarker value",
                "Tracks changes across multiple reports",
                "IMPROVED / WORSENED / STABLE status tags",
                "Side-by-side report comparison view",
            ],
        },
    ]

    pos = [(Inches(0.5), Inches(1.35)), (Inches(6.6), Inches(1.35)),
           (Inches(0.5), Inches(4.1)), (Inches(6.6), Inches(4.1))]

    for i, feat in enumerate(feature_data):
        lx, ty = pos[i]
        card(slide, lx, ty, Inches(5.85), Inches(2.5))
        # Icon + title
        rect(slide, lx, ty, Inches(5.85), Inches(0.5), GREEN if i % 2 == 0 else BLUE)
        txt(slide, f"{feat['icon']}  {feat['title']}",
            lx + Inches(0.15), ty + Inches(0.06), Inches(5.5), Inches(0.42),
            size=14, bold=True, color=WHITE)
        bullets(slide, feat["points"],
                lx + Inches(0.15), ty + Inches(0.6), Inches(5.5), Inches(1.8),
                size=12, color=DARK, gap=3)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 5 — KEY FEATURES (Module 2 — Doctors & Appointments)
# ════════════════════════════════════════════════════════════════════

def slide_05_features_doctors():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)
    slide_header(slide, "Key Features  |  Module 2", "Doctor Discovery & Appointment Booking")

    feature_data = [
        {
            "icon": "🩺",
            "title": "Symptom Intake & Specialty Suggestion",
            "color": GREEN,
            "points": [
                "Patient logs primary symptoms via structured form",
                "System matches symptoms to appropriate medical specialty",
                "Confidence score displayed alongside recommendation",
                "Emergency symptoms trigger urgent-care alert warning",
            ],
        },
        {
            "icon": "👨‍⚕️",
            "title": "Doctor Directory & Profiles",
            "color": BLUE,
            "points": [
                "Browse doctors filtered by recommended specialty",
                "Doctor cards: name, qualification, experience, fee, rating",
                "Full profile: bio, clinic address, languages spoken",
                "Star ratings and total review count displayed",
            ],
        },
        {
            "icon": "📅",
            "title": "Slot Selection & Booking",
            "color": GREEN,
            "points": [
                "Calendar date picker + time-slot grid per doctor",
                "Only available slots shown (booked slots greyed out)",
                "Attach a medical report to the consultation booking",
                "Booking confirmation screen with full summary",
            ],
        },
        {
            "icon": "🗂️",
            "title": "Appointments Management",
            "color": BLUE,
            "points": [
                "Full list of upcoming and past appointments",
                "Status badges: Confirmed / Completed / Cancelled",
                "Cancel appointment with reason selection",
                "Doctor contact info and clinic address on each card",
            ],
        },
    ]

    pos = [(Inches(0.5), Inches(1.35)), (Inches(6.6), Inches(1.35)),
           (Inches(0.5), Inches(4.1)),  (Inches(6.6), Inches(4.1))]

    for i, feat in enumerate(feature_data):
        lx, ty = pos[i]
        card(slide, lx, ty, Inches(5.85), Inches(2.5))
        rect(slide, lx, ty, Inches(5.85), Inches(0.5), feat["color"])
        txt(slide, f"{feat['icon']}  {feat['title']}",
            lx + Inches(0.15), ty + Inches(0.06), Inches(5.5), Inches(0.42),
            size=13, bold=True, color=WHITE)
        bullets(slide, feat["points"],
                lx + Inches(0.15), ty + Inches(0.6), Inches(5.5), Inches(1.8),
                size=12, color=DARK, gap=3)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 6 — APP ARCHITECTURE & DATA FLOW
# ════════════════════════════════════════════════════════════════════

def slide_06_architecture():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)
    slide_header(slide, "Technical Design", "Application Architecture & Data Flow")

    # ── Left: Layer diagram ──
    card(slide, Inches(0.5), Inches(1.35), Inches(5.5), Inches(5.55))
    txt(slide, "Layered Architecture",
        Inches(0.7), Inches(1.5), Inches(5.1), Inches(0.38),
        size=14, bold=True, color=GREEN)

    layers = [
        (GREEN,  "Presentation Layer",   "React Native Screens & Components (Expo)"),
        (BLUE,   "Navigation Layer",     "React Navigation: Stack + Bottom Tab Navigator"),
        (MID,    "State Layer",          "Zustand stores (Auth State, Report Wizard State)"),
        (GREEN,  "API / Service Layer",  "Axios API client — Calls FastAPI backend over HTTP"),
        (BLUE,   "Backend API Layer",    "FastAPI routers: Auth, Reports, Doctors, Appointments, AI"),
        (MID,    "Data Layer",           "PostgreSQL database via SQLAlchemy Async ORM"),
    ]

    y = Inches(2.02)
    for col, layer_name, layer_desc in layers:
        rect(slide, Inches(0.65), y, Inches(2.15), Inches(0.55), col)
        txt(slide, layer_name, Inches(0.68), y + Inches(0.1),
            Inches(2.1), Inches(0.38), size=10, bold=True, color=WHITE)
        txt(slide, layer_desc, Inches(2.92), y + Inches(0.1),
            Inches(2.9), Inches(0.38), size=10, color=DARK)
        if y < Inches(6.4):
            txt(slide, "▼", Inches(1.6), y + Inches(0.55),
                Inches(0.5), Inches(0.22), size=10, color=GREEN, align=PP_ALIGN.CENTER)
        y += Inches(0.78)

    # ── Right: User flow ──
    card(slide, Inches(6.4), Inches(1.35), Inches(6.45), Inches(5.55))
    txt(slide, "User Journey Flow",
        Inches(6.6), Inches(1.5), Inches(6.1), Inches(0.38),
        size=14, bold=True, color=BLUE)

    flow_steps = [
        ("1", "Register / Login", "User creates account or signs in with email & password"),
        ("2", "Upload Report", "Select lab PDF from gallery — backend parses biomarkers"),
        ("3", "View Analysis", "Read plain-English summary and flagged biomarker cards"),
        ("4", "Log Symptoms", "Describe primary complaint — app suggests medical specialty"),
        ("5", "Browse Doctors", "Filter doctor list by recommended specialty, fee, rating"),
        ("6", "Book Slot", "Pick date + time slot, attach report, confirm booking"),
        ("7", "Track & Manage", "View all appointments, cancel if needed, check status"),
    ]

    y = Inches(2.02)
    for num, step, desc in flow_steps:
        rect(slide, Inches(6.55), y, Inches(0.42), Inches(0.42),
             GREEN if int(num) % 2 == 1 else BLUE)
        txt(slide, num, Inches(6.55), y + Inches(0.04),
            Inches(0.42), Inches(0.34), size=12, bold=True,
            color=WHITE, align=PP_ALIGN.CENTER)
        txt(slide, step, Inches(7.08), y, Inches(2.5), Inches(0.28),
            size=12, bold=True, color=DARK)
        txt(slide, desc, Inches(7.08), y + Inches(0.27), Inches(5.6), Inches(0.28),
            size=10, color=MID)
        y += Inches(0.73)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 7 — CHALLENGES, LEARNINGS & FUTURE SCOPE
# ════════════════════════════════════════════════════════════════════

def slide_07_challenges():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)
    slide_header(slide, "Project Reflection", "Challenges, Learnings & Future Scope")

    # ── Challenges ──
    card(slide, Inches(0.5), Inches(1.35), Inches(3.9), Inches(5.55))
    rect(slide, Inches(0.5), Inches(1.35), Inches(3.9), Inches(0.5), GREEN)
    txt(slide, "Challenges Faced",
        Inches(0.65), Inches(1.42), Inches(3.65), Inches(0.4),
        size=14, bold=True, color=WHITE)
    challenges = [
        "Setting up Expo + NativeWind Tailwind CSS integration required careful babel config",
        "Managing async API calls and loading states across many screens",
        "Designing responsive layouts that work across different phone screen sizes",
        "Implementing JWT token refresh logic without logging out the user",
        "Coordinating navigation state between tab stacks and modal screens",
        "Handling file upload (PDF/image) from device with proper error handling",
    ]
    bullets(slide, challenges, Inches(0.65), Inches(1.98), Inches(3.65), Inches(4.7),
            size=11, color=DARK, gap=5)

    # ── Learnings ──
    card(slide, Inches(4.75), Inches(1.35), Inches(3.9), Inches(5.55))
    rect(slide, Inches(4.75), Inches(1.35), Inches(3.9), Inches(0.5), BLUE)
    txt(slide, "Key Learnings",
        Inches(4.9), Inches(1.42), Inches(3.65), Inches(0.4),
        size=14, bold=True, color=WHITE)
    learnings = [
        "Deep understanding of React Native component lifecycle and hooks",
        "How to structure a scalable project with screens, components, and API layers",
        "Using TypeScript interfaces to define strong data models for the entire app",
        "Integrating RESTful APIs with Axios and managing JWT authentication flows",
        "Using Zustand for lightweight global state instead of Redux boilerplate",
        "End-to-end mobile app development from design to deployment-ready build",
    ]
    bullets(slide, learnings, Inches(4.9), Inches(1.98), Inches(3.65), Inches(4.7),
            size=11, color=DARK, gap=5)

    # ── Future Scope ──
    card(slide, Inches(9.0), Inches(1.35), Inches(3.85), Inches(5.55))
    rect(slide, Inches(9.0), Inches(1.35), Inches(3.85), Inches(0.5), MID)
    txt(slide, "Future Scope",
        Inches(9.15), Inches(1.42), Inches(3.6), Inches(0.4),
        size=14, bold=True, color=WHITE)
    future = [
        "Push notification reminders for upcoming appointments",
        "Telemedicine / video consultation integration",
        "Offline report caching with SQLite for no-network access",
        "Google / Apple OAuth social login",
        "Multi-language support (Hindi, Gujarati, etc.)",
        "Wearable device integration (fitness trackers, smartwatches)",
        "Publish to Google Play Store & Apple App Store",
    ]
    bullets(slide, future, Inches(9.15), Inches(1.98), Inches(3.6), Inches(4.7),
            size=11, color=DARK, gap=5)


# ════════════════════════════════════════════════════════════════════
#  SLIDE 8 — THANK YOU / CONCLUSION
# ════════════════════════════════════════════════════════════════════

def slide_08_thankyou():
    slide = prs.slides.add_slide(BLANK)
    bg(slide)

    # Top green band
    rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.5), GREEN)
    txt(slide, "Thank You",
        Inches(0.5), Inches(0.2), Inches(12.3), Inches(1.1),
        size=60, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Project name
    txt(slide, "VitaLens  —  React Native Health & Doctor Booking App",
        Inches(0.5), Inches(1.6), Inches(12.3), Inches(0.5),
        size=18, color=DARK, align=PP_ALIGN.CENTER, bold=True)

    rect(slide, Inches(3.0), Inches(2.18), Inches(7.3), Pt(2), GREEN)

    # Summary card
    card(slide, Inches(0.5), Inches(2.35), Inches(12.35), Inches(2.3))
    txt(slide, "Project Summary",
        Inches(0.7), Inches(2.5), Inches(12), Inches(0.38),
        size=13, bold=True, color=GREEN)
    summary_pts = [
        "A fully functional cross-platform mobile application built using React Native (Expo) and TypeScript",
        "Covers complete user workflow: Register → Upload Report → View Analysis → Log Symptoms → Find Doctor → Book Appointment",
        "Backed by a FastAPI (Python) REST API with PostgreSQL database and JWT-based secure authentication",
    ]
    bullets(slide, summary_pts, Inches(0.7), Inches(2.95), Inches(12), Inches(1.55),
            size=12, color=DARK, gap=4)

    # Bottom two info cards
    # Left — Group
    card(slide, Inches(0.5), Inches(4.85), Inches(5.9), Inches(2.1))
    txt(slide, "Submitted By",
        Inches(0.7), Inches(5.0), Inches(5.5), Inches(0.35),
        size=13, bold=True, color=BLUE)
    txt(slide, "Uzair Shaikh  (22004500210246)",
        Inches(0.7), Inches(5.38), Inches(5.5), Inches(0.32),
        size=12, bold=True, color=DARK)

    # Right — Faculty
    card(slide, Inches(6.9), Inches(4.85), Inches(5.95), Inches(2.1))
    txt(slide, "Under the Guidance of",
        Inches(7.1), Inches(5.0), Inches(5.6), Inches(0.35),
        size=13, bold=True, color=GREEN)
    txt(slide, "Faculty / Guide:  Prof. Swati D. Londhe",
        Inches(7.1), Inches(5.38), Inches(5.6), Inches(0.32),
        size=12, bold=True, color=DARK)
    txt(slide, "Head of Department:  Prof. Monica Ma'am",
        Inches(7.1), Inches(5.72), Inches(5.6), Inches(0.32),
        size=12, bold=True, color=DARK)

    # Footer
    rect(slide, Inches(0), Inches(7.25), SLIDE_W, Inches(0.25), GREEN)
    txt(slide, "VitaLens  |  React Native Mobile App  |  University Project Presentation",
        Inches(0.4), Inches(7.25), Inches(12.5), Inches(0.25),
        size=9, color=WHITE, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════════════
#  BUILD & SAVE
# ════════════════════════════════════════════════════════════════════

print("Building VitaLens University Presentation...")
slide_01_title();   print("  [OK] Slide 1 - Title / Cover")
slide_02_problem(); print("  [OK] Slide 2 - Problem Statement & Objectives")
slide_03_tech_stack(); print("  [OK] Slide 3 - Technology Stack")
slide_04_features_reports(); print("  [OK] Slide 4 - Features: Reports & Home")
slide_05_features_doctors(); print("  [OK] Slide 5 - Features: Doctors & Appointments")
slide_06_architecture(); print("  [OK] Slide 6 - Architecture & Data Flow")
slide_07_challenges(); print("  [OK] Slide 7 - Challenges, Learnings & Future Scope")
slide_08_thankyou(); print("  [OK] Slide 8 - Thank You / Conclusion")

output = "VitaLens_University_Presentation.pptx"
prs.save(output)
print(f"\nSaved: {output}")
print("Open with Microsoft PowerPoint or Google Slides.")
